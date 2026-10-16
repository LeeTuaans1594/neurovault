"""Render fixture specs to real DOCX/PDF files + gold JSON.

Usage (from backend/ai_core):  .venv\\Scripts\\python.exe -m tests.fixtures.build_fixtures
Output is deterministic, so regenerated files only change when specs change.
"""

from __future__ import annotations

import json
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

try:
    from tests.fixtures.specs import ALL_SPECS
except ImportError:  # executed as a plain script
    from specs import ALL_SPECS  # type: ignore

HERE = Path(__file__).parent
DOCS_DIR = HERE / "docs"
GOLD_DIR = HERE / "gold"


# ───────────────────────────── DOCX ─────────────────────────────
def build_docx(spec: dict, out: Path) -> None:
    doc = Document()
    doc.core_properties.title = spec["title"]
    doc.add_heading(spec["title"], level=0)
    for sec in spec["sections"]:
        doc.add_heading(sec["heading"], level=sec["level"])
        for para in sec["body"]:
            doc.add_paragraph(para)
    tbl = spec.get("table")
    if tbl:
        t = doc.add_table(rows=1 + len(tbl["rows"]), cols=len(tbl["header"]))
        for j, h in enumerate(tbl["header"]):
            t.cell(0, j).text = h
        for i, row in enumerate(tbl["rows"], start=1):
            for j, val in enumerate(row):
                t.cell(i, j).text = val
    doc.save(str(out))


# ───────────────────────────── PDF ──────────────────────────────
W, H = 595, 842  # A4 points
MARGIN = 56
BODY, H1, H2 = 10.5, 18, 13


def _text(page, rect, text, size, bold=False) -> float:
    """Insert text in rect; return the height actually used."""
    rc = page.insert_textbox(fitz.Rect(*rect), text, fontsize=size,
                             fontname="hebo" if bold else "helv", align=0)
    if rc < 0:
        raise ValueError(f"Text overflow ({rc}) for: {text[:40]}")
    return (rect[3] - rect[1]) - rc


def _chrome(page, spec, page_no: int) -> None:
    page.insert_text((MARGIN, 32), spec["running_header"], fontsize=8, fontname="helv")
    page.insert_text((W / 2 - 10, H - 28), f"Page {page_no}", fontsize=8, fontname="helv")


def build_pdf(spec: dict, out: Path) -> None:
    doc = fitz.open()
    by_page: dict[int, list] = {}
    for sec in spec["sections"]:
        by_page.setdefault(sec["page"], []).append(sec)

    # Page 1: title + visible contents list
    p = doc.new_page(width=W, height=H)
    _chrome(p, spec, 1)
    y = 120
    y += _text(p, (MARGIN, y, W - MARGIN, y + 60), spec["title"], 24, bold=True) + 30
    y += _text(p, (MARGIN, y, W - MARGIN, y + 30), "Contents", H2, bold=True) + 8
    for sec in spec["sections"]:
        indent = 0 if sec["level"] == 1 else 18
        line = f"{sec['heading']} .......... {sec['page']}"
        y += _text(p, (MARGIN + indent, y, W - MARGIN, y + 20), line, BODY) + 2

    last_page = max(by_page)
    for page_no in range(2, last_page + 1):
        p = doc.new_page(width=W, height=H)
        _chrome(p, spec, page_no)
        y = 70
        secs = by_page.get(page_no, [])
        full = [s for s in secs if "column" not in s]
        cols = [s for s in secs if "column" in s]
        for sec in full:
            size = H1 if sec["level"] == 1 else H2
            y += _text(p, (MARGIN, y, W - MARGIN, y + 40), sec["heading"], size, bold=True) + 6
            for para in sec["body"]:
                y += _text(p, (MARGIN, y, W - MARGIN, y + 300), para, BODY) + 8
            y += 6
        if cols:
            gap, col_w = 24, (W - 2 * MARGIN - 24) / 2
            col_y = y
            for sec in cols:
                x0 = MARGIN if sec["column"] == "left" else MARGIN + col_w + gap
                cy = col_y
                cy += _text(p, (x0, cy, x0 + col_w, cy + 40), sec["heading"], H2, bold=True) + 6
                for para in sec["body"]:
                    cy += _text(p, (x0, cy, x0 + col_w, cy + 500), para, BODY) + 8
        tbl = spec.get("table")
        if tbl and tbl.get("page") == page_no:
            y += 10
            ncol = len(tbl["header"])
            cw = (W - 2 * MARGIN) / ncol
            for r, row in enumerate([tbl["header"], *tbl["rows"]]):
                for c, val in enumerate(row):
                    rect = fitz.Rect(MARGIN + c * cw, y, MARGIN + (c + 1) * cw, y + 20)
                    p.draw_rect(rect, width=0.5)
                    p.insert_text((rect.x0 + 4, rect.y0 + 14), val, fontsize=BODY,
                                  fontname="hebo" if r == 0 else "helv")
                y += 20

    doc.set_toc([[s["level"], s["heading"], s["page"]] for s in spec["sections"]])
    doc.set_metadata({"title": spec["title"], "author": "NeuroVault fixtures"})
    doc.save(str(out), deflate=True, garbage=3)
    doc.close()


# ───────────────────────────── gold ─────────────────────────────
def gold_for(spec: dict) -> dict:
    return {
        "name": spec["name"],
        "language": spec["language"],
        "file": f"{spec['name']}.{spec['format']}",
        "title": spec["title"],
        "headings": [
            {"text": s["heading"], "level": s["level"], "page": s.get("page")}
            for s in spec["sections"]
        ],
        "sections": [
            {"heading": s["heading"], "body": s["body"], "page": s.get("page")}
            for s in spec["sections"]
        ],
        "concepts": spec["gold_concepts"],
        "compounds": spec["gold_compounds"],
        "reading_order": spec.get("reading_order"),
        "noise_strings": [spec["running_header"]] if spec.get("running_header") else [],
    }


def build_all() -> list[Path]:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    GOLD_DIR.mkdir(parents=True, exist_ok=True)
    written = []
    for spec in ALL_SPECS:
        out = DOCS_DIR / f"{spec['name']}.{spec['format']}"
        (build_docx if spec["format"] == "docx" else build_pdf)(spec, out)
        gold_path = GOLD_DIR / f"{spec['name']}.json"
        gold_path.write_text(json.dumps(gold_for(spec), ensure_ascii=False, indent=2), encoding="utf-8")
        written += [out, gold_path]
    return written


if __name__ == "__main__":
    for p in build_all():
        print("wrote", p)
